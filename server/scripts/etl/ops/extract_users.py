# -*- coding: utf-8 -*-
r"""Operations (Canal Cities & Sinai) import, step 1 of 2 - the PEOPLE.

    python scripts/etl/ops/extract_users.py

Reads `Operations_Roster_TEST_DATA.xlsx` (sheet `Roster`, written and
self-checked by `BD\ECMS\Job Profiles\_working\generate_ops_roster.py`) and
writes `scripts/etl/data/ops/users.json`: one ECMS user document per person,
plus the unit -> manager references that go with them.

Unlike the BD/EC extract, the judgement part is NOT in this file: the roster
workbook already carries the unit id, the org level, the job profile id and the
reporting line, because the generator that invented these people decided all of
it in the open. What this step adds is the REFUSAL - it re-checks every one of
those columns against the live org chart and the eight loaded job profiles, so
a hand-edited or re-issued sheet cannot put somebody in a unit that does not
exist, on a rung its profile disagrees with, or under a manager who is not
there.

Checks:
  * every unit id is live (the ops org-unit load wins over the older export,
    so the `d-canal-ops` rename is respected) and is called what the sheet says;
  * the unit's structural TYPE is the right home for the org level
    (ASSISTANT_GENERAL->AGM, DEPARTMENT->DM, SECTION->SH/SP/JP/FR) - the same
    rule the employee form applies;
  * `jobProfileId` is one of the eight Operations profiles and its `orgLevel`
    matches the person's;
  * every manager is on the roster, sits on a higher rung, and every chain ends
    at the single person with no manager;
  * no duplicate employee number, Latin name or email, and no email already
    owned by one of the real BD / External-Contracts accounts.

THE ROSTER IS TEST DATA. Employee numbers are 90001+, a block nothing real uses,
so every document carries `isTestData: true` and the ids are `u-90001` ..
`u-90070` - one query removes the lot.
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "data")
OUT = os.path.join(DATA, "ops", "users.json")
LIVE_DEPTS = os.path.join(DATA, "departments.json")
OPS_DEPTS = os.path.join(DATA, "ops", "departments.json")
OPS_PROFILES = os.path.join(DATA, "ops", "jobProfiles.json")
BD_USERS = os.path.join(DATA, "bd-ec", "users.json")
SRC = os.environ.get(
    "OPS_ROSTER_XLSX",
    r"C:\Users\tariq\Desktop\work\Work laptop\BD\ECMS\Job Profiles"
    r"\1_Final_Deliverables\Operations_Canal_Sinai\Operations_Roster_TEST_DATA.xlsx",
)

# Which structural node type is the right home for each rung. A person sitting
# in the wrong type is what made the hand-moved BD people unsaveable.
HOME_TYPE = {
    "GM": "GENERAL",
    "AGM": "ASSISTANT_GENERAL",
    "DM": "DEPARTMENT",
    "SH": "SECTION",
    "SP": "SECTION",
    "JP": "SECTION",
    "FR": "SECTION",
}
# Seniority, most senior first - used to check a manager is actually above.
LADDER = ["CEO", "ACEO", "GM", "AGM", "DM", "SH", "SP", "JP", "FR"]
# The general department every one of these people belongs to.
GENERAL_DEPARTMENT = "g-canal"
# The invented block. Anything outside it would be claiming a real employee.
TEST_FIRST, TEST_LAST = 90001, 90099
COLUMNS = ["Employee No", "Arabic Name", "Name (Latin)", "Email", "Site", "Section",
           "Unit Id", "Unit Name", "Position", "Org Level", "Job Profile Id",
           "Reports To (No)", "Reports To (Name)"]


def load_units():
    """The live org chart, with the ops load layered on top (the rename wins)."""
    units = {}
    with open(LIVE_DEPTS, encoding="utf-8") as f:
        for d in json.load(f):
            units[d["id"]] = d["data"]
    with open(OPS_DEPTS, encoding="utf-8") as f:
        for d in json.load(f):
            units[d["id"]] = d["data"]
    return units


def read_roster(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    if "Roster" not in wb.sheetnames:
        sys.exit("REFUSING: %s has no `Roster` sheet" % os.path.basename(path))
    ws = wb["Roster"]
    head = next(ws.iter_rows(max_row=1, values_only=True))
    header = [str(c).strip() if c is not None else "" for c in head]
    if header != COLUMNS:
        sys.exit("REFUSING: the Roster sheet's columns changed:\n  expected %s\n  found    %s"
                 % (COLUMNS, header))

    rows = []
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[0] is None:
            continue
        rows.append(dict(
            row=n, emp=int(r[0]), arabic=str(r[1] or "").strip(),
            name=str(r[2] or "").strip(), email=str(r[3] or "").strip().lower(),
            site=str(r[4] or "").strip(), section=str(r[5] or "").strip(),
            unit=str(r[6] or "").strip(), unitName=str(r[7] or "").strip(),
            position=str(r[8] or "").strip(), level=str(r[9] or "").strip(),
            profile=str(r[10] or "").strip(),
            mgr=int(r[11]) if r[11] is not None else None,
            mgrName=str(r[12] or "").strip(),
        ))
    return rows


def check(rows, units, profiles, taken_emails):
    problems = []
    if not rows:
        problems.append("the Roster sheet has no people in it")

    by_emp, by_email, by_name = {}, {}, {}
    for row in rows:
        emp = row["emp"]
        at = "row %d (%s)" % (row["row"], emp)
        if emp in by_emp:
            problems.append("%s: employee number %d is used twice" % (at, emp))
        by_emp[emp] = row
        if not TEST_FIRST <= emp <= TEST_LAST:
            problems.append("%s: employee number is outside the invented %d-%d block"
                            % (at, TEST_FIRST, TEST_LAST))
        if not row["name"]:
            problems.append("%s: no Latin name" % at)
        if not row["arabic"]:
            problems.append("%s: no Arabic name" % at)
        by_name.setdefault(row["name"].lower(), []).append(emp)
        if not re.fullmatch(r"[a-z][a-z0-9.\-]*@eprom\.com", row["email"]):
            problems.append("%s: bad email %r" % (at, row["email"]))
        by_email.setdefault(row["email"], []).append(emp)
        if row["email"] in taken_emails:
            problems.append("%s: email %s already belongs to the real account %s"
                            % (at, row["email"], taken_emails[row["email"]]))

        unit = units.get(row["unit"])
        if unit is None:
            problems.append("%s: unit %s is not in the org chart" % (at, row["unit"]))
        else:
            if (unit.get("name") or "").strip() != row["unitName"]:
                problems.append("%s: unit %s is called %r, the sheet says %r"
                                % (at, row["unit"], unit.get("name"), row["unitName"]))
            want = HOME_TYPE.get(row["level"])
            if want is None:
                problems.append("%s: org level %r is not a rung" % (at, row["level"]))
            elif unit.get("type") != want:
                problems.append("%s: %s is a %s, which is not a home for a %s"
                                % (at, row["unit"], unit.get("type"), row["level"]))

        profile = profiles.get(row["profile"])
        if profile is None:
            problems.append("%s: job profile %s is not one of the Operations eight"
                            % (at, row["profile"]))
        elif profile.get("orgLevel") != row["level"]:
            problems.append("%s: org level %s but profile %s is %s"
                            % (at, row["level"], row["profile"], profile.get("orgLevel")))

    for email, owners in sorted(by_email.items()):
        if len(owners) > 1:
            problems.append("email %s is claimed by %s" % (email, owners))
    for name, owners in sorted(by_name.items()):
        if len(owners) > 1:
            problems.append("the name %r is used by %s" % (name, owners))

    # The reporting line: resolvable, senior, loop-free, one root.
    roots = [r for r in rows if r["mgr"] is None]
    for row in rows:
        at = "row %d (%s)" % (row["row"], row["emp"])
        if row["mgr"] is None:
            continue
        boss = by_emp.get(row["mgr"])
        if boss is None:
            problems.append("%s: reports to %d, who is not on the roster" % (at, row["mgr"]))
            continue
        if boss["name"] != row["mgrName"]:
            problems.append("%s: reports-to %d is %r, the sheet says %r"
                            % (at, row["mgr"], boss["name"], row["mgrName"]))
        if row["level"] in LADDER and boss["level"] in LADDER:
            if LADDER.index(boss["level"]) >= LADDER.index(row["level"]):
                problems.append("%s: manager %d is a %s, not senior to a %s"
                                % (at, row["mgr"], boss["level"], row["level"]))
    if len(roots) != 1:
        problems.append("expected exactly one person with no manager, found %d (%s)"
                        % (len(roots), ", ".join(str(r["emp"]) for r in roots)))
    for row in rows:
        seen, cur, looped = set(), row, False
        while cur is not None and cur["mgr"] is not None:
            if cur["emp"] in seen:
                problems.append("row %d (%s): the reporting chain loops"
                                % (row["row"], row["emp"]))
                looped = True
                break
            seen.add(cur["emp"])
            cur = by_emp.get(cur["mgr"])
        if not looped and cur is not None and roots and cur["emp"] != roots[0]["emp"]:
            problems.append("row %d (%s): the chain ends at %s, not at the top of the roster"
                            % (row["row"], row["emp"], cur["emp"]))
    return problems


def build(rows):
    users, dept_managers = [], {}
    for row in sorted(rows, key=lambda r: r["emp"]):
        uid = "u-%d" % row["emp"]
        doc = {
            "id": uid,
            "name": row["name"],
            "email": row["email"],
            "role": "EMPLOYEE",
            "status": "ACTIVE",
            "employeeId": row["emp"],
            "departmentId": row["unit"],
            "generalDepartmentId": GENERAL_DEPARTMENT,
            "orgLevel": row["level"],
            "jobProfileId": row["profile"],
            "location": row["site"],
            "projectName": "Strategic Tanks",
            # Traceability back to the roster - the app ignores these.
            "arabicName": row["arabic"],
            "sourceTitle": row["position"],
            # This roster is invented. Say so on every document.
            "isTestData": True,
        }
        if row["mgr"] is not None:
            doc["managerId"] = "u-%d" % row["mgr"]
        users.append(doc)
        # A person's own unit is theirs to run when they are its top rung: the
        # AGM runs the assistant-general unit, a DM their department, an SH
        # their section. That reference is what gives them a team in the app.
        if row["level"] in ("AGM", "DM", "SH"):
            if row["unit"] in dept_managers:
                sys.exit("REFUSING: %s would be run by both %s and %s"
                         % (row["unit"], dept_managers[row["unit"]], uid))
            dept_managers[row["unit"]] = uid
    return users, dept_managers


def main():
    units = load_units()
    with open(OPS_PROFILES, encoding="utf-8") as f:
        profiles = {p["id"]: p for p in json.load(f)}
    with open(BD_USERS, encoding="utf-8") as f:
        taken_emails = {u["email"].lower(): u["id"] for u in json.load(f)["users"]}

    rows = read_roster(SRC)
    problems = check(rows, units, profiles, taken_emails)
    if problems:
        print("REFUSING TO WRITE - %d problem(s):" % len(problems))
        for p in problems[:40]:
            print("  -", p)
        if len(problems) > 40:
            print("  ... and %d more" % (len(problems) - 40))
        sys.exit(1)

    users, dept_managers = build(rows)
    payload = {"users": users, "departmentManagers": dept_managers,
               "clearDepartmentManagers": []}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)

    by_level, by_site = {}, {}
    for u in users:
        by_level[u["orgLevel"]] = by_level.get(u["orgLevel"], 0) + 1
        by_site[u["location"]] = by_site.get(u["location"], 0) + 1
    print("wrote %d people -> %s" % (len(users), os.path.normpath(OUT)))
    print("  by rung   : " + ", ".join("%s %d" % (l, by_level[l]) for l in LADDER if l in by_level))
    print("  by site   : " + ", ".join("%s %d" % kv for kv in sorted(by_site.items())))
    print("  units run : %d (manager references written with the people)" % len(dept_managers))
    print("  ids       : %s .. %s   ALL TEST DATA" % (users[0]["id"], users[-1]["id"]))
    held = {u["jobProfileId"] for u in users}
    empty = sorted(p for p in profiles if p not in held)
    if empty:
        print("  no holder : " + ", ".join(empty) + "   (the GM seat is empty on purpose)")


if __name__ == "__main__":
    main()

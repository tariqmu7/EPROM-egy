# -*- coding: utf-8 -*-
r"""Operations import, step 1 of 2 - READ the skills this department needs.

The workbook `ECMS_Upload_1_SKILL.xlsx` is now the WHOLE company dictionary -
272 skills across four departments. This load is Operations only, so the sheet
is filtered down to the 117 skills the eight Operations profiles actually
require (read out of the JOB sheet, not typed here). Loading the other 155
would quietly import the Marketing department as a side effect.

21 of the 117 are shared with BD / External Contracts / Marketing and merge by
NAME onto one dictionary entry; their id therefore comes from the FIRST code in
the merged code cell (`BD-B-01 / EC-B-01 / MK-B-01 / OP-B-01` -> `sk-bd-b-01`),
which is the id those skills already carry in the database. The remaining 96 are
new (`sk-op-*`).

Field mapping and the wire shape are identical to `bd-ec/extract_skills.py` -
same five categories, same criticalities, same inline assessmentMethods block,
same ONE_TIME/ALL neutral default - so an Operations skill is indistinguishable
from a BD one.

    python extract_skills.py
"""
import json, os, re, sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = r"C:\Users\tariq\Desktop\work\Work laptop\BD\ECMS\Job Profiles\1_Final_Deliverables\ECMS_Import"
SRC = os.path.join(BASE, "ECMS_Upload_1_SKILL.xlsx")
JOBS = os.path.join(BASE, "ECMS_Upload_2_JOB.xlsx")
OUT = os.path.join(HERE, "..", "data", "ops", "skills.json")

OPS_UNIT = "Operations - Canal Cities & Sinai"

CATEGORIES = {"technical": "Technical", "safety": "Safety", "management": "Management",
              "soft skills": "Soft Skills", "behavioral": "Behavioral"}
CRITICALITIES = {"SAFETY_CRITICAL", "HIGH", "STANDARD", "LOW"}
METHODS = {"OJT_OBSERVATION", "WORK_RECORD_REVIEW", "WRITTEN_EXAM",
           "PRACTICAL_DEMO", "INTERVIEW", "THREE_SIXTY_EVALUATION"}

COL = {
    "name": "Name",
    "category": "Category (Technical/Safety/Management/Soft Skills/Behavioral)",
    "criticality": "Criticality (SAFETY_CRITICAL/HIGH/STANDARD/LOW)",
    "method": "Assessment Method (OJT_OBSERVATION/WORK_RECORD_REVIEW/WRITTEN_EXAM/PRACTICAL_DEMO/INTERVIEW)",
    "question": "Assessment Question",
    "link": "Assessment Link",
    "code": "Code",
    "description": "Description",
}


def skill_id(code, name):
    """`BD-B-01 / EC-B-01 / MK-B-01 / OP-B-01` -> `sk-bd-b-01`. Stable across re-runs."""
    first = (code or name).split("/")[0].strip()
    slug = re.sub(r"[^a-z0-9]+", "-", first.lower()).strip("-")
    return "sk-" + slug


def ops_skill_names():
    ws = openpyxl.load_workbook(JOBS, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    idx = {str(h).strip(): i for i, h in enumerate(rows[0])}
    names = []
    for row in rows[1:]:
        if str(row[idx["Department Name"]]).strip() == OPS_UNIT:
            names.append(str(row[idx["Skill Name"]]).strip())
    return names


def main():
    wanted = ops_skill_names()
    wanted_set = {n.lower() for n in wanted}
    if not wanted_set:
        sys.exit("no rows in the JOB sheet for %r" % OPS_UNIT)

    ws = openpyxl.load_workbook(SRC, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    for key, col in COL.items():
        if col not in idx:
            sys.exit("missing expected column: %s" % col)

    def cell(row, key):
        v = row[idx[COL[key]]]
        return "" if v is None else str(v).strip()

    skills, seen_ids, seen_names, problems = [], {}, {}, []
    for n, row in enumerate(rows[1:], start=2):
        name = cell(row, "name")
        if not name or name.lower() not in wanted_set:
            continue

        cat_raw = cell(row, "category").lower()
        if cat_raw not in CATEGORIES:
            problems.append("row %d: category %r is not one of the five ECMS categories" % (n, cat_raw))
        crit = cell(row, "criticality").upper()
        if crit not in CRITICALITIES:
            problems.append("row %d: criticality %r unknown" % (n, crit))
        method = cell(row, "method").upper()
        if method not in METHODS:
            problems.append("row %d: assessment method %r unknown" % (n, method))

        sid = skill_id(cell(row, "code"), name)
        if sid in seen_ids:
            problems.append("row %d: duplicate id %s (also row %d)" % (n, sid, seen_ids[sid]))
        seen_ids[sid] = n
        key = name.lower()
        if key in seen_names:
            problems.append("row %d: duplicate skill NAME %r (also row %d) - ECMS matches by name"
                            % (n, name, seen_names[key]))
        seen_names[key] = n

        levels = {}
        for lv in range(1, 6):
            desc = row[idx["Level %d Desc" % lv]]
            certs = row[idx["Level %d Certs" % lv]]
            levels[str(lv)] = {
                "level": lv,
                "description": "" if desc is None else str(desc).strip(),
                "requiredCertificates": [c.strip() for c in str(certs).split(",") if c.strip()] if certs else [],
            }
            if not levels[str(lv)]["description"]:
                problems.append("row %d: level %d has no description" % (n, lv))

        skills.append({
            "id": sid,
            "name": name,
            "category": CATEGORIES.get(cat_raw, "Technical"),
            "criticality": crit if crit in CRITICALITIES else "STANDARD",
            "levels": levels,
            "status": "APPROVED",
            "isArchived": False,
            "code": cell(row, "code"),
            "description": cell(row, "description"),
            "assessmentMethods": [{
                "id": "imp:%s" % sid,
                "method": method if method in METHODS else "OJT_OBSERVATION",
                "questions": [],
                "frequency": "ONE_TIME",
                "audience": "ALL",
            }],
            "assessmentMethod": method if method in METHODS else "OJT_OBSERVATION",
            "assessmentQuestion": cell(row, "question"),
            "assessmentLink": cell(row, "link"),
        })

    # Every requirement must land on a skill. A name the sheet cannot supply
    # would be a requirement that silently disappears from every gap figure.
    missing = sorted(wanted_set - set(seen_names))
    if missing:
        problems.append("%d skill name(s) required by the JOB sheet are not in the SKILL sheet: %s"
                        % (len(missing), ", ".join(missing[:5])))

    if problems:
        print("REFUSING TO WRITE - %d problem(s):" % len(problems))
        for p in problems[:40]:
            print("  -", p)
        sys.exit(1)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(skills, f, ensure_ascii=False, indent=1)

    shared = [s for s in skills if not s["id"].startswith("sk-op-")]
    print("wrote %d skills -> %s" % (len(skills), os.path.normpath(OUT)))
    print("  new to Operations : %d (sk-op-*)" % (len(skills) - len(shared)))
    print("  shared by name    : %d (already in the dictionary under BD/EC ids)" % len(shared))
    by_cat, by_crit = {}, {}
    for s in skills:
        by_cat[s["category"]] = by_cat.get(s["category"], 0) + 1
        by_crit[s["criticality"]] = by_crit.get(s["criticality"], 0) + 1
    print("  by category       : " + ", ".join("%s %d" % kv for kv in sorted(by_cat.items())))
    print("  by criticality    : " + ", ".join("%s %d" % kv for kv in sorted(by_crit.items())))


if __name__ == "__main__":
    main()
